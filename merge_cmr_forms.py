
import os

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(HERE, "CMR_merged.xlsx")


KEY_CELLS = [(1, 0), (7, 0), (7, 3), (13, 0), (17, 0), (25, 0), (25, 1), (25, 2), (25, 5),
             (50, 0), (50, 2), (53, 2), (58, 0), (58, 2)]


BORDER_STYLES = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick",
                 6: "double", 7: "hair", 8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot",
                 11: "dashDotDot", 12: "mediumDashDotDot", 13: "slantDashDot"}




OLD_GEORGIAN_FONTS = ("acadnusx", "litnusx")
NEW_GEORGIAN_FONT = "Sylfaen"      
GEORGIAN_FONT = {
    "a": "ა", "b": "ბ", "g": "გ", "d": "დ", "e": "ე", "v": "ვ", "z": "ზ", "T": "თ",
    "i": "ი", "k": "კ", "l": "ლ", "m": "მ", "n": "ნ", "o": "ო", "p": "პ", "J": "ჟ",
    "r": "რ", "s": "ს", "t": "ტ", "u": "უ", "f": "ფ", "q": "ქ", "R": "ღ", "y": "ყ",
    "S": "შ", "C": "ჩ", "c": "ც", "Z": "ძ", "w": "წ", "W": "ჭ", "x": "ხ", "j": "ჯ",
    "h": "ჰ",
}


def to_georgian(text):
  
    return "".join(GEORGIAN_FONT.get(ch, ch) for ch in text)


H_ALIGN = {0: None, 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous"}
V_ALIGN = {0: "top", 1: "center", 2: "bottom", 3: "justify"}


def find_file(start, end):
    for name in os.listdir(HERE):
        if name.lower().startswith(start.lower()) and name.lower().endswith(end.lower()):
            return os.path.join(HERE, name)
    raise FileNotFoundError(f"No file starting with '{start}' and ending with '{end}' in {HERE}")



def fingerprint(sheet):
    
    parts = []
    for r, c in KEY_CELLS:
        v = sheet.cell_value(r, c) if r < sheet.nrows and c < sheet.ncols else ""
        parts.append(" ".join(str(v).split()))
    return "|".join(parts)


def is_empty(sheet):
    return sheet.nrows == 0 or all(not str(v).strip() for r in range(sheet.nrows) for v in sheet.row_values(r))



class StyleCopier:
    

    def __init__(self, book):
        self.book = book
        self.cache = {}

    def color(self, index):
        rgb = self.book.colour_map.get(index)
        return "FF{:02X}{:02X}{:02X}".format(*rgb) if rgb else None

    def side(self, line, color_index):
        style = BORDER_STYLES.get(line)
        if not style:
            return Side()
        return Side(style=style, color=self.color(color_index))

    def styles(self, xf_index):
        if xf_index in self.cache:
            return self.cache[xf_index]
        xf = self.book.xf_list[xf_index]
        f = self.book.font_list[xf.font_index]
        old_font = f.name.lower().replace(" ", "") in OLD_GEORGIAN_FONTS
        look = dict(size=f.height / 20, bold=bool(f.bold), italic=bool(f.italic),
                    underline="single" if f.underline_type else None,
                    color=self.color(f.colour_index) or "FF000000")  
        font = Font(name=NEW_GEORGIAN_FONT if old_font else f.name, **look)
        georgian_font = Font(name=NEW_GEORGIAN_FONT, **look)        


        b = xf.border
        border = Border(left=self.side(b.left_line_style, b.left_colour_index),
                        right=self.side(b.right_line_style, b.right_colour_index),
                        top=self.side(b.top_line_style, b.top_colour_index),
                        bottom=self.side(b.bottom_line_style, b.bottom_colour_index))

        bg = xf.background
        fill_color = self.color(bg.pattern_colour_index)
        fill = PatternFill("solid", fgColor=fill_color) if bg.fill_pattern == 1 and fill_color else PatternFill()

        a = xf.alignment
        align = Alignment(horizontal=H_ALIGN.get(a.hor_align), vertical=V_ALIGN.get(a.vert_align, "bottom"),
                          wrap_text=bool(a.text_wrapped))

        number_format = self.book.format_map[xf.format_key].format_str if xf.format_key in self.book.format_map else "General"
        result = (font, georgian_font, border, fill, align, number_format, old_font)
        self.cache[xf_index] = result
        return result


def is_old_font(book, font_index):
    return book.font_list[font_index].name.lower().replace(" ", "") in OLD_GEORGIAN_FONTS


def convert_text(book, src, r, c, text, cell_font_is_old):

   
    runs = src.rich_text_runlist_map.get((r, c))
    if not runs:
        return (to_georgian(text), True) if cell_font_is_old else (text, False)

    pieces, any_old = [], False
    starts = [0] + [offset for offset, _ in runs]
    old_flags = [cell_font_is_old] + [is_old_font(book, font_index) for _, font_index in runs]
    for i, start in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(text)
        piece = text[start:end]
        if old_flags[i]:
            piece, any_old = to_georgian(piece), True
        pieces.append(piece)
    return "".join(pieces), any_old


def copy_sheet(book, src, dst, copier):
    
    for r in range(src.nrows):
        for c in range(src.ncols):
            cell = src.cell(r, c)
            out = dst.cell(row=r + 1, column=c + 1)
            font, georgian_font, border, fill, align, number_format, old_font = copier.styles(cell.xf_index)
            if cell.ctype == xlrd.XL_CELL_DATE:
                out.value = xlrd.xldate_as_datetime(cell.value, book.datemode)
            elif cell.ctype == xlrd.XL_CELL_TEXT:
                out.value, has_old_part = convert_text(book, src, r, c, cell.value, old_font)
                if has_old_part:
                    font = georgian_font                        
            elif cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_BOOLEAN):
                out.value = cell.value
            out.font, out.border, out.fill, out.alignment = font, border, fill, align
            out.number_format = number_format

    for rlo, rhi, clo, chi in src.merged_cells:      
        dst.merge_cells(start_row=rlo + 1, end_row=rhi, start_column=clo + 1, end_column=chi)

    for c, info in src.colinfo_map.items():              
        dst.column_dimensions[get_column_letter(c + 1)].width = info.width / 256
    for r, info in src.rowinfo_map.items():             
        dst.row_dimensions[r + 1].height = info.height / 20



def main():
    new_path = find_file("IEL", "2024.xls")
    old_path = find_file("11", ".xls")

    out = Workbook()
    out.remove(out.active)
    seen, names = set(), set()
    stats = {"copied": 0, "duplicates": 0, "empty": 0}


    for path, label in [(new_path, ""), (old_path, " (old)")]:
        print(f"Reading {os.path.basename(path)} with formatting ...")
        book = xlrd.open_workbook(path, formatting_info=True)
        copier = StyleCopier(book)
        for sheet in book.sheets():
            if is_empty(sheet):
                stats["empty"] += 1
                continue
            key = fingerprint(sheet)
            if key in seen:
                stats["duplicates"] += 1
                continue
            seen.add(key)

            name = (sheet.name + label)[:31]               
            base, n = name, 2
            while name in names:                           
                name = f"{base[:27]} #{n}"
                n += 1
            names.add(name)

            copy_sheet(book, sheet, out.create_sheet(name), copier)
            stats["copied"] += 1
            if stats["copied"] % 250 == 0:
                print(f"  {stats['copied']:,} CMRs copied ...")

    print("Saving (this can take a minute) ...")
    out.save(OUTPUT)
    print(f"\nUnique CMRs (sheets):  {stats['copied']:,}")
    print(f"Duplicates skipped:    {stats['duplicates']:,}")
    print(f"Empty sheets skipped:  {stats['empty']:,}")
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
